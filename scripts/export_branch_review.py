"""Branch review export for manual correction (SU10S-18). READ-ONLY.

    python3 -m scripts.export_branch_review [--out ~/branch_review.xlsx]

Population: stores that are SERVING (hold prices), LIVE (db/query.py
live_store_clause - loaded within 3 days of their chain's latest load) and
PHYSICAL. One row per store; a store can carry several issues:

    NO_CITY            city_canonical NULL
    NO_ADDRESS         effective address NULL / empty
    PLACEHOLDER        effective address is a placeholder ("unknown", "לא ידוע", "-" ...)
    NO_HOUSE_NUMBER    effective address has no digit
    GEOCODE_REJECTED   Nominatim answered, but no candidate matched the input
                       street (+ house number), or the match is > 25 km from the
                       CBS city centroid
    GEOCODE_CITY_MISMATCH  the street matched, but Nominatim puts it in another
                       city than city_canonical - usually a wrong city (SU10S-25)
    GEOCODE_NO_MATCH   Nominatim returned nothing for the address (not in the brief;
                       same "check this address" signal, so surfaced too)
    NO_COORDINATES     no lat/lon at all, not even a city centroid (the Sunday
                       geo_centroids / geo_nominatim run has not reached it yet)
    MAYBE_ONLINE       name/address suggests online / delivery / fulfilment

"Effective address" = COALESCE(NULLIF(btrim(address_override), ''), address),
the same expression scripts/geo_nominatim.py geocodes.

GEOCODE_* ARE REPLAYED, NOT STORED: the geocoder persists only accepted
results. Every answer it ever got is in its on-disk cache, and its accept rules
(geo_nominatim.evaluate(), SU10S-25) are deterministic, so this replays them
against the cache. It NEVER calls Nominatim - an uncached query is reported as
"not geocoded yet" instead. A hand-placed row (geo_source = 'manual') is
never geocoded, so it never carries a GEOCODE_* issue.

Sheets, in order (Dude's triage, SU10S-18 follow-up):
    "עדיפות"                   worth manual effort: GEOCODE_REJECTED,
                               GEOCODE_CITY_MISMATCH, NO_CITY. issue_detail says what
                               Nominatim matched - reason, road/number/city, km from
                               the CBS city centroid.
    "לבדיקה"                   optional: everything else (NO_MATCH, NO_HOUSE_NUMBER,
                               PLACEHOLDER, ...).
    "bulk-awaiting-StoresFull" NO_ADDRESS rows of chains that publish NO address for
                               any branch. They wait for StoresFull ingestion, not
                               hand entry. Which chains is computed, not listed: a
                               chain qualifies when none of its stores has a non-empty
                               FEED `address` (the loader-owned column, not the
                               override). At SU10S-18: שופרסל, ויקטורי, קינג סטור,
                               שפע ברכת השם, שוק העיר, חצי חינם.
A store is on each sheet at most once. A no-address-chain store with another
issue is on the bulk sheet for NO_ADDRESS and on עדיפות / לבדיקה for the rest.

The four trailing columns are for the reviewer; the apply session reads them:
    correct_city     -> scraper/city_names.py STORE_CITY_OVERRIDES (durable)
    correct_address  -> stores.address_override
    is_physical=no   -> stores.is_physical = false
then a targeted re-geocode.
"""
from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import text

from db.db import connect
from db.query import live_store_clause
from scripts.geo_centroids import load_cbs_centroids
from scripts.geo_nominatim import (
    CACHE_PATH,
    EFFECTIVE_ADDRESS_SQL,
    CityResolver,
    build_geo_input,
    cache_key,
    evaluate,
)

BULK_LABEL = "BULK — awaiting StoresFull ingestion"
SHEET_PRIORITY = "עדיפות"
SHEET_OPTIONAL = "לבדיקה"
SHEET_BULK = "bulk-awaiting-StoresFull"
PRIORITY_ISSUES = {"GEOCODE_REJECTED", "GEOCODE_CITY_MISMATCH", "NO_CITY"}

# Chains whose feed carries no address for ANY store (see module doc).
_NO_ADDRESS_CHAINS_SQL = """
    SELECT s.chain_id, c.name
    FROM stores s LEFT JOIN chains c ON c.chain_id = s.chain_id
    GROUP BY s.chain_id, c.name
    HAVING count(*) FILTER (WHERE btrim(coalesce(s.address, '')) <> '') = 0
    ORDER BY c.name
"""

ISSUE_ORDER = ["NO_CITY", "NO_ADDRESS", "PLACEHOLDER", "NO_HOUSE_NUMBER",
               "GEOCODE_CITY_MISMATCH", "GEOCODE_REJECTED", "GEOCODE_NO_MATCH",
               "NO_COORDINATES", "MAYBE_ONLINE"]

_PLACEHOLDERS = {"unknown", "none", "null", "n/a", "na", "-", "--", "---", ".", "0",
                 "לא ידוע", "לא ידועה", "אין", "אין כתובת", "לא קיים", "כללי", "?"}

_ONLINE_RE = re.compile(
    r"אונליין|online|אינטרנט|internet|משלוח|delivery|פיק\s*-?\s*אפ|pick\s*-?\s*up|"
    r"מחסן|לוגיסטי|מרלו\"?ג|fulfil|e-?commerce|דיגיטל|digital|וירטואל|virtual|"
    r"הזמנות|סיטונאות|wholesale",
    re.IGNORECASE,
)

COLUMNS = ["stores.id", "chain", "store_id", "store_name", "city_canonical", "raw city",
           "address", "geo_precision", "issues", "issue_detail",
           "correct_address", "correct_city", "is_physical", "notes"]
REVIEWER_COLS = {"correct_address", "correct_city", "is_physical", "notes"}

_POP_SQL = """
    SELECT s.id, s.chain_id, c.name AS chain, s.store_id, s.store_name,
           s.city_canonical, s.city, {eff} AS address, s.geo_precision,
           s.lat, s.lon, s.geo_source
    FROM stores s
    LEFT JOIN chains c ON c.chain_id = s.chain_id
    WHERE s.is_physical
      AND EXISTS (SELECT 1 FROM prices p WHERE p.store_fk = s.id)
      {live}
    ORDER BY c.name, s.store_id
"""


def load_cache() -> dict:
    try:
        return json.loads(Path(CACHE_PATH).read_text(encoding="utf-8"))
    except Exception:
        return {}


def _cache_get(cache: dict, params: dict):
    """Exactly the geocoder's cache key; None = never asked."""
    return cache.get(cache_key(params))


class Replay:
    """What evaluate() needs besides the row: the cache, CBS centroids, cities."""

    def __init__(self) -> None:
        self.cache = load_cache()
        self.centroids = load_cbs_centroids()
        self.cities = CityResolver(list(self.centroids))


def replay_geocode(r: dict, rp: Replay) -> tuple[str | None, str]:
    """(issue or None, detail) for a row the geocoder would target."""
    cache = rp.cache
    addr, city, geo_input = build_geo_input(r["address"], r["city_canonical"])
    hits = _cache_get(cache, {"street": addr, "city": city})
    if hits is None:
        return None, "not geocoded yet (next scrp-geocode run)"
    if not hits:
        hits = _cache_get(cache, {"q": f"{addr}, {city}"})
        if hits is None:
            return None, "not geocoded yet (next scrp-geocode run)"
    if not hits:
        return "GEOCODE_NO_MATCH", f"no Nominatim result for '{geo_input}'"
    centroid = rp.centroids.get(city)
    if centroid is None and r["lat"] is not None:
        centroid = (r["lat"], r["lon"])
    res = evaluate(hits, addr, city, centroid, r["chain"], rp.cities)
    if res["precision"]:
        return None, ""
    # What Nominatim matched, for the reviewer: "NO_MATCH: הרצל 12, רמת גן 3.1 km".
    where = ", ".join(str(v) for v in (" ".join(str(x) for x in (res["road"], res["house_number"]) if x),
                                       res["city"]) if v)
    dist = f" {res['dist_km']:.1f} km" if res["dist_km"] is not None else ""
    what = f"{res['reason']}: {where}{dist}"
    return ("GEOCODE_CITY_MISMATCH" if res["reason"] == "CITY_MISMATCH" else "GEOCODE_REJECTED"), what


def geocodable(address: str | None, city: str | None) -> bool:
    """Mirror of geo_nominatim._TARGETS_SQL's address/city filters."""
    a = (address or "").strip()
    return (bool(a) and a.lower() not in {"unknown", "none", "null", "n/a", "-"}
            and bool(re.search(r"[0-9]", a)) and bool((city or "").strip()))


def issues_for(r: dict, rp: Replay) -> tuple[list[str], list[str]]:
    issues, detail = [], []
    addr = (r["address"] or "").strip()
    if not r["city_canonical"]:
        issues.append("NO_CITY")
    if not addr:
        issues.append("NO_ADDRESS")
    elif addr.lower() in _PLACEHOLDERS:
        issues.append("PLACEHOLDER")
        detail.append(f"address = '{addr}'")
    elif not re.search(r"[0-9]", addr):
        issues.append("NO_HOUSE_NUMBER")
        if addr in {(r["city"] or "").strip(), (r["city_canonical"] or "").strip()}:
            detail.append("address is just the city name")
    if r["lat"] is None:
        issues.append("NO_COORDINATES")
    if r["geo_source"] != "manual" and geocodable(r["address"], r["city_canonical"]):
        g, why = replay_geocode(r, rp)
        if g:
            issues.append(g)
        if why:
            detail.append(why)
    blob = f"{r['store_name'] or ''} {addr}"
    m = _ONLINE_RE.search(blob)
    if m:
        issues.append("MAYBE_ONLINE")
        detail.append(f"matched '{m.group(0)}'")
    return issues, detail


def to_row(r: dict, issues: list[str], detail: list[str]) -> list:
    # The last four are the reviewer's columns and must start empty.
    return [r["id"], r["chain"], r["store_id"], r["store_name"], r["city_canonical"],
            r["city"], r["address"], r["geo_precision"], ", ".join(issues),
            "; ".join(detail), None, None, None, None]


def write_sheet(ws, rows: list[list]) -> None:
    ws.sheet_view.rightToLeft = True
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"stores.id": 8, "chain": 13, "store_id": 9, "store_name": 30, "city_canonical": 15,
              "raw city": 15, "address": 32, "geo_precision": 11, "issues": 30, "issue_detail": 45,
              "correct_address": 30, "correct_city": 16, "is_physical": 11, "notes": 30}
    fill = PatternFill("solid", fgColor="FFF6D5")
    for idx, name in enumerate(COLUMNS, 1):
        letter = ws.cell(row=1, column=idx).column_letter
        ws.column_dimensions[letter].width = widths[name]
        if name in REVIEWER_COLS:
            for (cell,) in ws.iter_rows(min_row=1, max_row=max(ws.max_row, 2), min_col=idx, max_col=idx):
                cell.fill = fill
    dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    ws.add_data_validation(dv)
    col = ws.cell(row=1, column=COLUMNS.index("is_physical") + 1).column_letter
    dv.add(f"{col}2:{col}{max(ws.max_row, 2)}")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def main() -> None:
    ap = argparse.ArgumentParser(description="Export the branch review workbook (read-only).")
    ap.add_argument("--out", default=str(Path.home() / "branch_review.xlsx"))
    args = ap.parse_args()

    conn = connect()
    try:
        live_sql, live_params = live_store_clause(conn)
        rows = conn.execute(text(_POP_SQL.format(eff=EFFECTIVE_ADDRESS_SQL, live=live_sql)),
                            live_params).mappings().all()
        no_addr = {r[0]: r[1] for r in conn.execute(text(_NO_ADDRESS_CHAINS_SQL)).all()}
    finally:
        conn.rollback()
        conn.close()

    rp = Replay()
    sheets: dict[str, list] = {SHEET_PRIORITY: [], SHEET_OPTIONAL: [], SHEET_BULK: []}
    cat = Counter()
    per_sheet_chain: dict[str, Counter] = {k: Counter() for k in sheets}
    per_sheet_issue: dict[str, Counter] = {k: Counter() for k in sheets}
    for r in rows:
        issues, detail = issues_for(r, rp)
        for i in issues:
            cat[i] += 1
        if r["chain_id"] in no_addr and "NO_ADDRESS" in issues:
            sheets[SHEET_BULK].append((r["chain"] or "", "", r["store_id"], to_row(r, ["NO_ADDRESS"], [BULK_LABEL])))
            per_sheet_chain[SHEET_BULK][r["chain"]] += 1
            per_sheet_issue[SHEET_BULK]["NO_ADDRESS"] += 1
            issues = [i for i in issues if i != "NO_ADDRESS"]
        if not issues:
            continue
        issues.sort(key=ISSUE_ORDER.index)
        sheet = SHEET_PRIORITY if PRIORITY_ISSUES & set(issues) else SHEET_OPTIONAL
        sheets[sheet].append((r["chain"] or "", ISSUE_ORDER.index(issues[0]), r["store_id"],
                              to_row(r, issues, detail)))
        per_sheet_chain[sheet][r["chain"]] += 1
        for i in issues:
            per_sheet_issue[sheet][i] += 1

    wb = Workbook()
    wb.remove(wb.active)
    for name, entries in sheets.items():
        entries.sort(key=lambda t: t[:3])
        write_sheet(wb.create_sheet(name), [t[3] for t in entries])
    wb.save(args.out)

    print(f"population (serving + live + physical): {len(rows)}")
    print(f"no-address chains (computed): {', '.join(no_addr.values())}")
    print("stores per issue (whole population, before the split):")
    for i in ISSUE_ORDER:
        print(f"  {i:18} {cat[i]}")
    for name in sheets:
        print(f"sheet {name}: {len(sheets[name])} rows")
        print("   by issue: " + ", ".join(f"{i} {per_sheet_issue[name][i]}" for i in ISSUE_ORDER
                                          if per_sheet_issue[name][i]))
        print("   by chain: " + ", ".join(f"{ch} {n}" for ch, n in
                                          sorted(per_sheet_chain[name].items(), key=lambda kv: -kv[1])))
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
