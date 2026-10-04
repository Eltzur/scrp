"""Manual-pin sheet: stores still without an exact position (SU11A-15). READ-ONLY.

    python3 -m scripts.export_pin_sheet [--out ~/pin_sheet.xlsx]

Population: LIVE (db/query.py live_store_clause) PHYSICAL stores whose
position is only a city centroid (geo_precision = 'city') or missing
(lat IS NULL). One row per store, sorted by chain, city, store id.

Columns: store_fk, chain, store_id, store_name, feed_address (the effective
address the geocoder sees: address_override, else address), city_canonical,
current_precision, google_maps_url (a search for chain + branch + address +
city), and lat / lon left BLANK for Dude to fill from Google Maps (right-click
the pin -> the first line is "lat, lon"). Filled rows become manual pins
(geo_source = 'manual'), which no geocoder run ever overwrites (SU10S-25).

Same pattern as scripts/export_branch_review.py: one SELECT, rolled back,
nothing written to the database.
"""
from __future__ import annotations

import argparse
import urllib.parse
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import text

from db.db import connect
from db.query import live_store_clause

# The geocoder's effective address (scripts/geo_nominatim.EFFECTIVE_ADDRESS_SQL), aliased.
_EFFECTIVE_S = "COALESCE(NULLIF(btrim(s.address_override), ''), s.address)"
COLUMNS = ["store_fk", "chain", "store_id", "store_name", "feed_address", "city_canonical",
           "current_precision", "google_maps_url", "lat", "lon"]
FILL_COLS = {"lat", "lon"}
WIDTHS = {"store_fk": 9, "chain": 14, "store_id": 9, "store_name": 32, "feed_address": 32,
          "city_canonical": 16, "current_precision": 12, "google_maps_url": 40, "lat": 12, "lon": 12}

_POP_SQL = """
    SELECT s.id, c.name AS chain, s.store_id, s.store_name,
           {eff} AS address, s.city_canonical, s.geo_precision, s.lat
    FROM stores s LEFT JOIN chains c ON c.chain_id = s.chain_id
    WHERE s.is_physical
      AND (s.geo_precision = 'city' OR s.lat IS NULL)
      {live}
    ORDER BY c.name, s.city_canonical NULLS LAST, s.store_id
"""


def maps_url(chain: str | None, name: str | None, address: str | None, city: str | None) -> str:
    q = " ".join(p for p in (chain, name, address, city) if p)
    return "https://www.google.com/maps/search/?api=1&query=" + urllib.parse.quote(q)


def main() -> None:
    ap = argparse.ArgumentParser(description="Export the manual-pin sheet (read-only).")
    ap.add_argument("--out", default=str(Path.home() / "pin_sheet.xlsx"))
    args = ap.parse_args()

    conn = connect()
    try:
        live_sql, live_params = live_store_clause(conn)
        rows = conn.execute(text(_POP_SQL.format(eff=_EFFECTIVE_S, live=live_sql)),
                            live_params).mappings().all()
    finally:
        conn.rollback()
        conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "pins"
    ws.sheet_view.rightToLeft = True
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    per_chain: Counter = Counter()
    for r in rows:
        precision = r["geo_precision"] if r["lat"] is not None else "none"
        per_chain[r["chain"]] += 1
        ws.append([r["id"], r["chain"], r["store_id"], r["store_name"], r["address"],
                   r["city_canonical"], precision,
                   maps_url(r["chain"], r["store_name"], r["address"], r["city_canonical"]),
                   None, None])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    fill = PatternFill("solid", fgColor="FFF6D5")
    for idx, name in enumerate(COLUMNS, 1):
        letter = ws.cell(row=1, column=idx).column_letter
        ws.column_dimensions[letter].width = WIDTHS[name]
        if name in FILL_COLS:
            for (cell,) in ws.iter_rows(min_row=1, max_row=max(ws.max_row, 2), min_col=idx, max_col=idx):
                cell.fill = fill
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    wb.save(args.out)

    print(f"stores without an exact position (live, physical): {len(rows)}")
    for chain, n in per_chain.most_common():
        print(f"  {chain:<16}{n:>5}")
    print(f"\nsheet: {args.out}")


if __name__ == "__main__":
    main()
